import streamlit as st
import openpyxl

st.set_page_config(page_title="تحديث الشيت الذكي", page_icon="🚗")
st.title("🚗 تطبيق تحديث الشيت الذكي")

# دالة لتنظيف اللوحة من أي مسافات لضمان المطابقة 100%
def clean_plate(plate_str):
    if not plate_str: return ""
    return str(plate_str).replace(" ", "").strip()

if 'step' not in st.session_state:
    st.session_state.step = 1

uploaded_main = st.file_uploader("1. ارفع الشيت الأساسي (المجمع)", type=["xlsx"])
uploaded_port = st.file_uploader("2. ارفع شيت المحفظة الجديدة", type=["xlsx"])

if st.button("🔄 إعادة البدء من جديد"):
    for key in list(st.session_state.keys()):
        del st.session_state[key]
    st.session_state.step = 1
    st.rerun()
    
if uploaded_main and uploaded_port:
    if st.session_state.step == 1:
        if st.button("🔍 تحليل المحفظة ومطابقة اللوحات"):
            try:
                wb_main = openpyxl.load_workbook(uploaded_main)
                sheet_main = wb_main.active
                wb_port = openpyxl.load_workbook(uploaded_port, data_only=True)
                sheet_port = wb_port.active

                port_data = []
                port_plates = set()
                
                # قراءة المحفظة (بدءاً من الصف 2)
                for row in range(2, sheet_port.max_row + 1):
                    plate = sheet_port.cell(row=row, column=1).value
                    if plate:
                        cleaned_plate = clean_plate(plate)
                        car_type = sheet_port.cell(row=row, column=4).value
                        port_data.append({'original_plate': plate, 'cleaned_plate': cleaned_plate, 'type': car_type})
                        port_plates.add(cleaned_plate)

                current_owners = set()
                plates_found = 0
                
                # قراءة الشيت الأساسي والمطابقة بدون مسافات
                for row in range(2, sheet_main.max_row + 1):
                    main_plate = sheet_main.cell(row=row, column=1).value
                    if main_plate:
                        cleaned_main_plate = clean_plate(main_plate)
                        if cleaned_main_plate in port_plates:
                            owner = sheet_main.cell(row=row, column=3).value
                            if owner:
                                current_owners.add(str(owner).strip())
                            plates_found += 1

                st.session_state.port_data = port_data
                st.session_state.port_plates = port_plates
                st.session_state.current_owners = list(current_owners)
                st.session_state.plates_found = plates_found
                st.session_state.port_plates_count = len(port_data)
                st.session_state.step = 2
                st.rerun()
            except Exception as e:
                st.error(f"حدث خطأ أثناء التحليل: {e}")

    elif st.session_state.step == 2:
        st.info(f"📊 **نتيجة التحليل:** تم العثور على {st.session_state.plates_found} لوحة مكررة في الشيت الأساسي من أصل {st.session_state.port_plates_count} لوحة مرفوعة في المحفظة.")
        
        if st.session_state.current_owners:
            owners_str = "، ".join(st.session_state.current_owners)
            st.warning(f"⚠️ اللوحات المكررة مسجلة حالياً بأسماء الملاك التالين: **{owners_str}**")
            default_name = st.session_state.current_owners[0]
        else:
            st.success("✨ هذه اللوحات جديدة تماماً وغير مسجلة مسبقاً في الشيت.")
            default_name = ""

        new_owner = st.text_input("✍️ أدخل اسم المالك الذي تريد تسجيل المحفظة باسمه:", value=default_name)
        
        if st.button("✅ تأكيد الاسم والتجهيز"):
            if new_owner:
                st.session_state.new_owner = new_owner
                st.session_state.step = 3
                st.rerun()
            else:
                st.error("يرجى إدخال اسم المالك!")

    elif st.session_state.step == 3:
        st.success(f"جاهز لتحديث الشيت وتسجيل المحفظة باسم: **{st.session_state.new_owner}**")
        
        if st.button("🚀 استخراج وتنزيل الملف النهائي"):
            try:
                wb_main = openpyxl.load_workbook(uploaded_main)
                sheet_main = wb_main.active

                # مسح القديم
                for row in range(sheet_main.max_row, 1, -1):
                    main_plate = sheet_main.cell(row=row, column=1).value
                    if main_plate:
                        cleaned_main_plate = clean_plate(main_plate)
                        if cleaned_main_plate in st.session_state.port_plates:
                            sheet_main.delete_rows(row, 1)

                # إضافة الجديد مع الحفاظ على شكل اللوحة الأصلي
                for data in st.session_state.port_data:
                    sheet_main.append([data['original_plate'], data['type'], st.session_state.new_owner])

                output_file = "Updated_Master_Final.xlsx"
                wb_main.save(output_file)

                with open(output_file, "rb") as f:
                    st.download_button("📥 تحميل ملف الإكسل الآن", f, file_name="Updated_Master_Final.xlsx")
            except Exception as e:
                st.error(f"حدث خطأ أثناء الاستخراج: {e}")
